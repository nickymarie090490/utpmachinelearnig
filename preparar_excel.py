"""Reconstruye los datos desde el Excel original; conserva solo columnas necesarias."""
from pathlib import Path
import random, json
import openpyxl
import pandas as pd
from texto import TEMAS, limpiar, etiquetas_reglas, SEMILLA
P=Path(__file__).parent
fuente=P.parent/'upload/Etiquetados.xlsx'
if not fuente.exists(): fuente=P/'Etiquetados.xlsx'
w=openpyxl.load_workbook(fuente,read_only=True,data_only=True)
s=w.active
rows=s.iter_rows(values_only=True); header=next(rows)
manual=[]
for row in rows:
 d=dict(zip(header,row))
 if all(d[t] in (0,1) for t in TEMAS):
  manual.append({k:d[k] for k in ['listing_id','review_id','comments']+TEMAS})
m=pd.DataFrame(manual); assert len(m)>0
m.to_csv(P/'datos/muestra_manual.csv',index=False)
groups=set(m.listing_id); texts=set(m.comments.map(limpiar)); seen=set(); sample=[]
rng=random.Random(SEMILLA); count=0; total=0
for row in s.iter_rows(min_row=2,values_only=True):
 d=dict(zip(header,row)); total+=1
 if d['listing_id'] in groups: continue
 text=d['comments']; clean=limpiar(text or '')
 if clean in texts or clean in seen or not any(c.isalpha() for c in clean): continue
 seen.add(clean); count+=1
 item={k:d[k] for k in ['listing_id','review_id','comments']}
 if len(sample)<25000: sample.append(item)
 else:
  j=rng.randrange(count)
  if j<25000: sample[j]=item
v=pd.DataFrame(sample); v[TEMAS]=[etiquetas_reglas(t) for t in v.comments]
v.to_csv(P/'datos/entrenamiento_debil.csv.gz',index=False,compression='gzip')
(P/'datos/origen.json').write_text(json.dumps({'originales':total,'manuales':len(m),'entrenamiento':len(v),'elegibles':count},indent=2))
print('Datos preparados',total,len(m),len(v),flush=True)
