"""Genera informe y bitácora Word a partir de resultados finales, sin entrenar."""
from pathlib import Path
import json,textwrap
import pandas as pd
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
P=Path(__file__).parent; O=P.parent/'entrega_final';O.mkdir(exist_ok=True)
def read(n):return json.loads((P/'datos'/n).read_text())
def csv(n):return pd.read_csv(P/'datos'/n)
e=read('evidencia_sobreajuste.json');metrics=csv('metricas.csv');tests=read('pruebas_resumen.json');origin=read('origen.json')

def base():
 d=Document();s=d.sections[0];s.page_height=Inches(11.7);s.page_width=Inches(8.3);s.top_margin=Inches(.7);s.bottom_margin=Inches(.65);s.left_margin=s.right_margin=Inches(.75)
 for name in ['Normal','Body Text']:
  st=d.styles[name];st.font.name='Calibri';st.font.size=Pt(10.5);st.paragraph_format.space_after=Pt(6);st.paragraph_format.line_spacing=1.08
 for name,size in [('Title',23),('Heading 1',16),('Heading 2',12),('Heading 3',11)]:
  st=d.styles[name];st.font.name='Calibri';st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string('000000' if name=='Title' else '17365D');st.paragraph_format.space_before=Pt(8);st.paragraph_format.space_after=Pt(7)
 d.styles['Caption'].font.size=Pt(9);d.styles['Caption'].font.color.rgb=RGBColor.from_string('475569')
 f=s.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT;f.add_run('Equipo 6  |  ')
 field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)
 
 for st in d.styles:
  for border in st.element.xpath('.//w:pBdr'):border.getparent().remove(border)
 d.core_properties.author='Equipo 6';d.core_properties.title='Clasificación multietiqueta de reseñas de alojamientos';return d

def para(d,text):return d.add_paragraph(text)
def head(d,text,level=1):
 p=d.add_heading(text,level)
 if getattr(d,'_next_page',False):p.paragraph_format.page_break_before=True;d._next_page=False
def page(d):d._next_page=True
def table(d,caption,headers,rows,widths=None):
 d.add_paragraph(caption,'Caption');t=d.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 if widths:
  for col,w in zip(t.columns,widths):col.width=Inches(w)
 for cell,v in zip(t.rows[0].cells,headers):cell.text=str(v)
 for row in rows:
  for cell,v in zip(t.add_row().cells,row):cell.text=str(v)
 for i,row in enumerate(t.rows):
  trpr=row._tr.get_or_add_trPr()
  cant=OxmlElement('w:cantSplit');trpr.append(cant)
  if i==0:repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
  for j,cell in enumerate(row.cells):
   cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   if widths:cell.width=Inches(widths[j])
   prop=cell._tc.get_or_add_tcPr();shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'17365D' if i==0 else ('EFF4F9' if i%2==0 else 'FFFFFF'));prop.append(shade)
   borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    b=OxmlElement('w:'+edge);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
   prop.append(borders);marg=OxmlElement('w:tcMar')
   for ed in ['top','left','bottom','right']:
    n=OxmlElement('w:'+ed);n.set(qn('w:w'),'85');n.set(qn('w:type'),'dxa');marg.append(n)
   prop.append(marg)
   for p in cell.paragraphs:
    p.paragraph_format.space_after=Pt(1);p.paragraph_format.space_before=Pt(1);p.paragraph_format.line_spacing=1
    for run in p.runs:
     run.font.size=Pt(9);run.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '111827');run.bold=(i==0)
 empty=d.add_paragraph();empty.paragraph_format.space_after=Pt(1);empty.paragraph_format.line_spacing=Pt(2);empty.add_run().font.size=Pt(1)
 return t

def image(d,file,caption,width=6.7):
 d.add_picture(str(P/'datos'/file),width=Inches(width));d.add_paragraph(caption,'Caption')

def pipeline():
 fig,ax=plt.subplots(figsize=(10,7));ax.set_xlim(0,10);ax.set_ylim(-.3,10);ax.axis('off')
 def box(x,y,w,h,text,color='#edf2f7'):
  ax.add_patch(FancyBboxPatch((x,y),w,h,boxstyle='round,pad=0.08',facecolor=color,edgecolor='#64748b'));ax.text(x+w/2,y+h/2,text,ha='center',va='center',fontsize=10)
 def arr(a,b):ax.annotate('',xy=b,xytext=a,arrowprops={'arrowstyle':'->','color':'#334155','lw':1.4})
 box(2.2,8.9,5.6,.7,'Excel original y criterios de etiquetado');box(2.2,7.6,5.6,.8,'Excluir test humano, copias y textos vacíos\nMuestrear 25.000 reseñas y crear etiquetas débiles');arr((5,8.9),(5,8.4))
 box(2.2,6.3,5.6,.7,'PARTICIÓN POR ALOJAMIENTO antes del ajuste');arr((5,7.6),(5,7))
 box(.2,4.85,4.2,.8,'Entrenamiento 19.931 / 6.440 alojamientos\nValidación 5.069 / 1.610 alojamientos');box(6,4.85,3.8,.8,'Test humano reservado\n100 reseñas / 90 alojamientos');arr((4,6.3),(2.3,5.65));arr((6.7,6.3),(7.9,5.65))
 box(.2,2.6,4.2,1.6,'PIPELINE CLÁSICO\nLimpieza y stemming → TF-IDF\nColumnTransformer + log-longitud\nGridSearchCV GroupKFold → LR / SVM','#e9f1ff');box(5.2,2.6,4.6,1.6,'PIPELINE KERAS\nLimpieza → TF-IDF → SVD256\nColumnTransformer + longitud → escala\nMLP128/64 + dropout + early stopping','#fff3df');arr((1.5,4.85),(1.5,4.2));arr((3.7,4.85),(6.7,4.2))
 box(2.2,1.1,5.6,.8,'Evaluación final común y costos\nCurvas, errores, importancia e idiomas');arr((2.3,2.6),(3.5,1.9));arr((7.5,2.6),(6.5,1.9));arr((9.5,4.85),(9.5,1.5));arr((9.5,1.5),(7.8,1.5))
 box(2.2,.0,5.6,.6,'Streamlit con regresión logística y explicación local');arr((5,1.1),(5,.6));fig.tight_layout();fig.savefig(P/'datos/pipeline.png',dpi=180);plt.close(fig)

def informe():
 pipeline();d=base();d.add_paragraph('Clasificación multietiqueta de reseñas de alojamientos en Santiago','Title')
 para(d,'Informe del proyecto final de Machine Learning\nMaestría en Analítica de Datos · Universidad Tecnológica de Panamá\nGrupo A · Equipo 6 · Profesor Juan Montenegro\nRevisión final del 4 de octubre de 2026')
 head(d,'Introducción')
 para(d,'Elegimos la regresión logística para la aplicación de operaciones. Su Macro F1 de 0,913 queda por debajo del 0,934 de SVM, pero ofrece puntuaciones por tema, una explicación lineal y un modelo de 0,297 MB. La red Keras mejoró respecto a la versión anterior, aunque alcanzó 0,781 y siguió fallando en muchas menciones de ruido. La comparación, y no la victoria de la red, sustenta la decisión.')
 para(d,'El Problema 2 pide identificar temas en texto real de alojamientos. El informe sigue los nueve requerimientos técnicos y termina con las seis reflexiones de la guía. Todas las cifras pertenecen a la ejecución final; las pruebas unitarias verifican contratos de código, mientras que la muestra manual evalúa clasificación. La fecha del dump no consta y la muestra humana ya se había consultado en versiones previas. Estas dos limitaciones impiden afirmar cumplimiento pleno de procedencia y test ciego.')
 head(d,'1 Formulación del problema')
 para(d,'Decisión concreta: el área de operaciones utilizará las menciones de limpieza, ruido, ubicación, anfitrión y precio para priorizar qué reseñas revisar y qué temas investigar por alojamiento.')
 para(d,'La unidad es una reseña de la población suministrada en Etiquetados.xlsx, identificada como Santiago en el notebook. Cada salida vale 1 cuando el texto menciona el tema, incluso si lo elogia, lo niega o lo describe; vale 0 cuando no lo menciona. Las etiquetas no son excluyentes. No se pronostica un evento futuro: se clasifica el texto ingresado, sin ventana temporal de predicción. No se infiere una queja solo por detectar un tema.')
 para(d,'Mantuvimos Macro F1 como métrica principal, definida desde el planteamiento anterior. Da el mismo peso a las cinco categorías y combina precisión y recall: un falso negativo deja una mención sin revisar y un falso positivo consume tiempo de operaciones. Ruido y precio son escasos en entrenamiento; exactitud por etiqueta podría ocultar esos errores. No contamos con costos monetarios para fijar una ponderación económica diferente.')
 table(d,'Tabla 1 Referencia trivial y métricas complementarias',['Modelo trivial','Macro F1','Micro F1','Coincidencia completa'],[['Mayoritaria por etiqueta','0,140','0,390','2%']], [2.4,1.25,1.25,1.7])
 para(d,'La referencia aprende las mayorías solo en entrenamiento y siempre marca ubicación, con las otras etiquetas en cero. La Tabla 1 establece el piso; Micro F1 y coincidencia completa complementan el criterio principal.')
 # Página 2
 page(d);head(d,'1 Formulación del problema y diccionario de datos')
 table(d,'Tabla 2 Diccionario de columnas usadas y descartadas',['Campo','Significado y unidad','Uso'],[
 ['listing_id','Identificador del alojamiento; categórico sin unidad','Agrupar particiones y remuestreo'],['review_id','Identificador de reseña; sin unidad','Trazabilidad local'],['comments','Texto de la reseña; cadena','Entrada del modelo'],['limpieza ruido ubicación anfitrión precio','Cinco indicadores binarios 0 o 1','Targets humanos o débiles'],['log_longitud','log(1 + número de caracteres)','Variable estructurada derivada'],['date','Fecha de reseña; día calendario','Auditoría del período, no predictor'],['seleccion_manual','Marca original de selección','No se usa para construir el target; etiquetas completas determinan muestra'],['reviewer_id y reviewer_name','Identificador y nombre de reseñador','Descartados; no son predictores'],['Campos desconocidos','No hay campos usados sin significado definido','Fecha de dump ausente']], [1.5,3.1,2.0])
 head(d,'Construcción de etiquetas',2)
 para(d,'Se preservaron las 100 etiquetas completas de la hoja Etiquetado. El criterio distingue aseo y suciedad, sonido y silencio, localización y cercanía, atención y comunicación del anfitrión, y costo o valor. Por ejemplo, un apartamento silencioso menciona ruido y una ubicación excelente menciona ubicación. Las anotaciones existentes son humanas, dirigidas y sin documentación de acuerdo entre anotadores; no afirmamos doble revisión.')
 para(d,'Las 25.000 reseñas de desarrollo se etiquetaron con reglas multilingües. Sirven como supervisión débil, no como verdad de referencia. Se corrigió la regla inglesa cost para no etiquetar Costanera como precio, y se incluyeron loud y loudly para sonido. La normalización conserva texto original localmente; elimina HTML, acentos y puntuación para el procesamiento, aplica minúsculas, stopwords españolas e inglesas y stemming español.')
 table(d,'Tabla 3 Calidad y versión del archivo original',['Registro','Resultado'],[['Filas originales','690.112'],['Excluidas por alojamiento del test','13.555'],['Copias del texto manual','21.748'],['Sin letras o vacías','601'],['Duplicados normalizados','54.100'],['Elegibles después de exclusiones','600.108'],['Período de reseñas','13 nov 2010 a 1 jul 2026'],['Fecha del dump','No registrada; debe acreditarse con la descarga original']], [3.6,3.0])
 para(d,'La Tabla 3 procede del recorrido completo del Excel. Se conservó un hash SHA-256 en origen.json para identificar exactamente la versión recibida. Ni las fechas de reseña ni la modificación del Excel sustituyen la fecha del volcado exigida por la guía.')
 # Página3
 page(d);head(d,'2 Partición honesta de los datos')
 para(d,'Reservamos los alojamientos de las 100 reseñas humanas y excluimos sus textos normalizados antes de muestrear desarrollo. GroupShuffleSplit separa 80% y 20% de entidades automáticas. La población independiente se cuenta por alojamiento, porque varias reseñas pueden compartir anfitrión, entorno y vocabulario. No usamos KFold como validador final ni una partición por filas.')
 table(d,'Tabla 4 Particiones de la ejecución final',['Conjunto','Reseñas','Alojamientos','Etiqueta'],[['Entrenamiento','19.931','6.440','Automática'],['Validación','5.069','1.610','Automática'],['Evaluación final','100','90','Humana dirigida']], [2.0,1.2,1.6,1.8])
 para(d,'No hay alojamientos compartidos ni copias de textos humanos en desarrollo. Los pipelines ajustan vocabulario, IDF, SVD y escaladores solo con entrenamiento. GridSearchCV recibe tres pliegues GroupKFold; cada candidato reajusta la preparación dentro de su pliegue. Keras ajusta su preparación en entrenamiento y usa la validación automática para seleccionar candidato y detener épocas.')
 cv=csv('particiones.csv');agg=cv.groupby(['modelo','validador']).agg({'Macro F1 automático':'mean','grupos_compartidos':'mean'})
 rows=[]
 for n in ['LogisticRegression','LinearSVC']:
  na=agg.loc[(n,'KFold ingenuo')];ok=agg.loc[(n,'GroupKFold correcto')]
  rows.append([n,f"{na.iloc[0]:.6f}",f"{ok.iloc[0]:.6f}",f"{(na.iloc[0]-ok.iloc[0]):+.6f}",'2.690 / 0'])
 table(d,'Tabla 5 Partición ingenua y correcta con el mismo candidato',['Modelo','F1 ingenuo','F1 agrupado','Diferencia','Grupos compartidos'],rows,[1.9,1.05,1.05,1.0,1.6])
 para(d,'La Tabla 5 compara tres pliegues del candidato clásico seleccionado, con parámetros fijos y preprocesamiento reajustado. La regresión logística tiene una diferencia de -0,002794; SVM prácticamente no cambia. El KFold ingenuo comparte en promedio 2.690 alojamientos por pliegue. La fuga potencial no siempre eleva la métrica: una diferencia pequeña no legitima compartir entidades. Además, la validación automática premia semejanza con reglas y no estima calidad humana.')
 head(d,'Límite de la evaluación final',2)
 para(d,'Las etiquetas humanas se consultaron únicamente después de seleccionar modelos en esta ejecución. Sin embargo, esas mismas 100 reseñas ya se usaron en versiones anteriores y el fallo de ruido se conocía. Por ello, no son un test completamente ciego. La selección de esta versión utilizó validación automática, pero no se puede borrar el conocimiento previo. Se requiere otra muestra humana de alojamientos independientes para cumplir estrictamente la reserva final y estimar generalización.')
 para(d,'La estructura temporal no exige TimeSeriesSplit en este planteamiento descriptivo. Si se pretendiera anticipar reseñas futuras, habría que separar también por fecha. Las fechas no entran como predictor, y no se mezcla información de listings o calendar porque el Problema 2 requiere comments de reviews.')
 # Página4
 page(d);head(d,'3 Nivel 1 Modelo clásico')
 para(d,'Entrenamos LogisticRegression y LinearSVC, ambos con OneVsRestClassifier y cinco decisiones binarias. La limpieza y TF-IDF forman parte de la preparación dentro del pipeline. ColumnTransformer combina el texto con log(1 + longitud), y StandardScaler normaliza la longitud usando estadísticas del entrenamiento.')
 table(d,'Tabla 6 Configuración y búsqueda de los clásicos',['Elemento','Valor y razón'],[['TF-IDF','ngram_range candidatos (1,1) y (1,2); palabras individuales o pares'],['min_df','2; descarta términos que aparecen en un solo texto de entrenamiento'],['max_features','6.000; limita vocabulario y costo de memoria'],['Tokenización','Stopwords y stemming español; raíces, no comprensión semántica'],['Regresión logística','liblinear, class_weight balanced, max_iter 1.000'],['SVM lineal','dual auto, class_weight balanced, max_iter 3.000'],['Optimización','GridSearchCV del pipeline completo, C 0,5 o 2, cuatro candidatos por clasificador'],['Validación','GroupKFold tres pliegues, shuffle y semilla fija'],['Selección final','Ambos: C 2 y unigramas; sin optimizar umbral humano']], [1.7,4.9])
 para(d,'Los unigramas ganaron en la búsqueda, por lo que no se añadió complejidad de bigramas al modelo final. La selección usó Macro F1 automático: 0,917 para logística y 0,961 para SVM. En el conjunto automático separado, los valores fueron 0,935 y 0,971. La mayor afinidad del SVM con etiquetas de reglas no demuestra igual rendimiento en lenguaje libre.')
 head(d,'Salidas y costos de error',2)
 para(d,'Logística ofrece una salida sigmoid por tema. El umbral fijo 0,5 convierte esa puntuación en detección. Las puntuaciones no se calibraron y no deben presentarse como certeza de que una persona se queja. SVM ofrece márgenes y decisiones; no se transformaron esos márgenes en porcentajes inventados.')
 para(d,'class_weight balanced compensa la escasez por clase en cada ajuste. En el entrenamiento automático, las prevalencias son limpieza 22,5%, ruido 4,0%, ubicación 54,1%, anfitrión 43,3% y precio 4,5%. La muestra manual tiene una distribución diferente: incluye 25 menciones de ruido y 33 de precio entre 100 reseñas. Esa selección dirigida limita extrapolar al flujo real de operaciones.')
 para(d,'Un falso positivo de precio puede surgir al reducir preciosa y precio a la raíz preci. Ese error observado se conserva y se analiza; no se añadieron reglas de predicción para forzar etiquetas sobre el test. Corregir representaciones lingüísticas más finas exigiría volver a validar con otra muestra independiente.')
 # Página5
 page(d);head(d,'4 Suficiencia de los datos')
 para(d,'La Figura 1 cuenta alojamientos independientes. Se entrenó con subconjuntos anidados del 10%, 25%, 50%, 75% y 100% de los 6.440 alojamientos de entrenamiento, manteniendo fijas las 5.069 reseñas de validación. Cada subconjunto reajusta su preparación; Keras aplica early stopping con la misma configuración seleccionada. Las métricas comparan contra etiquetas automáticas.')
 image(d,'curva_aprendizaje.png','Figura 1 Curvas de aprendizaje por número de alojamientos',6.6)
 curve=csv('curva_aprendizaje.csv');rows=[]
 for frac in [.1,.25,.5,.75,1.]:
  v=curve[curve.proporcion==frac];rows.append([f'{frac:.0%}',str(v.grupos.iloc[0]),str(v.filas.iloc[0]),*[f"{v[v.modelo==n]['Macro F1 automático'].iloc[0]:.3f}" for n in ['LogisticRegression','LinearSVC','MLP']]])
 table(d,'Tabla 7 Tamaños y desempeño en validación fija',['Porción','Grupos','Filas','Logística','SVM','Keras'],rows,[.8,1.0,1.0,1.2,1.2,1.4])
 para(d,'Logística sube de 0,927 a 0,935 entre 75% y 100%; SVM de 0,969 a 0,971. Ambos siguen ascendiendo al 100%, por lo que declaramos insuficiencia para afirmar que ya no aportarían más datos, como exige la guía. Keras es inestable: alcanza 0,788 al 50%, cae a 0,771 al 75% y sube a 0,781 al 100%. Tampoco permite afirmar una meseta estable.')
 para(d,'El volumen permite entrenar la arquitectura escogida en CPU, pero muchas filas no equivalen a información independiente o etiquetas confiables. Más datos débiles pueden reforzar errores de las reglas. La siguiente mejora debería aumentar alojamientos y etiquetas humanas diversas, no solo duplicar reseñas. Esta curva se usó como análisis de suficiencia; no se volvió a elegir configuración con sus resultados.')
 # Página6
 page(d);head(d,'5 Nivel 2 Red neuronal en Keras')
 para(d,'La representación densa obligatoria se obtiene con TruncatedSVD sobre TF-IDF. Se probaron 64 componentes sin ponderación, 128 con ponderación y 256 con ponderación; sus Macro F1 automáticos fueron 0,639, 0,728 y 0,781. Se eligió el tercero por validación, sin consultar el test durante esa selección. Cambiaron compresión y ponderación conjuntamente: no atribuimos toda la mejora a una sola causa.')
 table(d,'Tabla 8 Arquitectura final de Keras',['Elemento','Configuración'],[['Entrada','256 componentes SVD + longitud escalada = 257 números'],['Capas Dense','128 ReLU, 64 ReLU y 5 sigmoid; tres capas Dense'],['Dropout','0,20 después de cada capa oculta; desactivado en inferencia'],['Optimizador','Adam con learning_rate 0,001'],['Pérdida y batch','BCE ponderada por etiqueta; batch 128'],['Pesos positivos','Raíz de negativos/positivos en entrenamiento: 1,85; 4,89; 0,92; 1,15; 4,58'],['Early stopping','val_loss, patience 6, min_delta 0,0001, restore_best_weights True'],['Épocas','12 de máximo 100; pesos restaurados de la época 6']], [1.5,5.1])
 image(d,'curvas_perdida.png','Figura 2 Pérdidas reales con controles y sin dropout ni early stopping',6.6)
 para(d,'Los 256 componentes conservan 53,1% de la varianza TF-IDF de entrenamiento. La preparación completa se ajusta después de particionar. La pérdida ponderada aumenta el costo de omitir etiquetas escasas; no usa prevalencias del test. La salida sigmoid permite varias detecciones simultáneas.')
 para(d,'La parada se activó realmente. La comparación sin controles conserva datos, arquitectura, ponderación, semilla y Adam, pero entrena 100 épocas sin dropout ni parada. Los pesos exportados son los restaurados; la inferencia carga Keras con compile=False y verifica hashes del modelo y preparador.')
 # Página7
 page(d);head(d,'5 Tratamiento del sobreajuste')
 table(d,'Tabla 9 Pérdidas comparables en inferencia sin dropout',['Medida','Con controles','Sin ambos controles'],[['BCE entrenamiento',f"{e['loss_train_restaurada']:.6f}",f"{e['loss_train_sin_control']:.6f}"],['BCE validación',f"{e['loss_val_restaurada']:.6f}",f"{e['loss_val_sin_control']:.6f}"],['Brecha validación menos entrenamiento',f"{e['loss_val_restaurada']-e['loss_train_restaurada']:.6f}",f"{e['loss_val_sin_control']-e['loss_train_sin_control']:.6f}"],['Macro F1 automático','0,781','0,783']], [3.2,1.7,1.7])
 reduction=(e['loss_val_sin_control']-e['loss_val_restaurada'])/e['loss_val_sin_control']*100
 para(d,f'La Tabla 9 muestra una reducción de {reduction:.1f}% de BCE de validación frente a la ablación. Sin controles, el entrenamiento queda casi perfectamente ajustado y la validación empeora: evidencia compatible con sobreajuste. Dropout y early stopping ayudan conjuntamente a contener esa pérdida, pero no se aislaron sus efectos. El F1 de la ablación es ligeramente mayor; minimizar pérdida no equivale a maximizar F1 con umbral fijo. Se usó una semilla, por lo que la evidencia es exploratoria.')
 para(d,'La Figura 2 muestra la pérdida ponderada usada durante fit, con dropout activo solo al entrenar. La Tabla 9 usa BCE no ponderada y training=False para comparar ambos conjuntos bajo las mismas condiciones. No se restaron directamente pérdidas que se midieron de manera diferente.')
 head(d,'6 Comparación de los dos niveles')
 rows=[]
 for _,r in metrics.iterrows():rows.append([r.modelo,f"{r['Macro F1 manual']:.3f}",f"{r['Micro F1 manual']:.3f}",f"{r['Coincidencia completa']:.0%}",f"{r.entrenamiento_s:.3f}",f"{r.inferencia_ms_resena:.3f}",f"{r.tamano_MB:.3f}"])
 table(d,'Tabla 10 Comparación sobre las mismas 100 reseñas humanas',['Modelo','Macro F1','Micro F1','Exacta','Entrenam s','Infer ms','MB'],rows,[1.8,.8,.8,.7,.85,.85,.8])
 para(d,'El entrenamiento incluye búsqueda y ajuste final; excluye curvas, ablación e interpretabilidad. La inferencia es la mediana de siete repeticiones en lote de 100, tras calentamiento, sin tiempo de cargar archivos o TensorFlow. Los costos se midieron en CPU x86_64 con hilos limitados y no son benchmarks del servidor público. SVM y logística probaron cuatro candidatos con tres pliegues; Keras tres candidatos con validación fija. La igualdad de presupuesto de boosting no aplica al Problema 2 de texto.')
 para(d,'Los intervalos exploratorios al 95%, remuestreando 90 alojamientos 500 veces, son 0,882-0,945 para logística, 0,900-0,961 para SVM y 0,730-0,821 para Keras. No eliminan el sesgo de selección de la muestra ni prueban diferencias causales. La red queda 0,132 puntos de F1 por debajo de logística y 0,153 por debajo de SVM.')
 # Página8
 page(d);head(d,'6 Matrices de confusión y análisis de errores')
 image(d,'matrices_confusion.png','Figura 3 Matrices por tema y modelo con filas reales y columnas predichas',6.6)
 para(d,'La Figura 3 usa una matriz 2 × 2 por etiqueta, porque una única matriz de cinco clases sería incorrecta en multietiqueta. En ruido, Keras tiene 9 verdaderos positivos y 16 falsos negativos; logística 24 y 1. En precio, logística tiene 27 verdaderos positivos, 6 falsos negativos y 6 falsos positivos. Los 26 errores completos de logística y los 49 de Keras coinciden en 24 reseñas; 2 fallan solo en logística y 25 solo en la red. SVM falla en 20 y todas esas reseñas también fallan en Keras.')
 table(d,'Tabla 11 Cinco casos erróneos de logística o Keras resumidos sin identidades',['Caso','Referencia y predicción','Dificultad observada'],[
 ['Vista preciosa y bien ubicado','Manual ubicación; los tres añaden precio','Stemming acerca preciosa a precio'],['Suciedad con ruido de TV y anfitrión atento','Manual cuatro temas; Keras omite limpieza','Reseña larga con varias señales competidoras'],['Metro y ayuda de la anfitriona','Manual ubicación y anfitrión; Keras solo ubicación','La atención se expresa indirectamente'],['Silencioso en pleno centro','Manual ruido y ubicación; Keras solo ubicación','Silencio es mención del tema, no ausencia'],['Maltrato y acusación de robo','Manual anfitrión; logística y Keras añaden limpieza y precio','Lenguaje mixto y palabras asociadas fuera de su contexto']], [1.5,2.6,2.5])
 para(d,'La Tabla 11 describe los primeros cinco errores conjuntos en el orden fijo del archivo, no casos escogidos para favorecer un modelo. Se parafrasearon textos y eliminaron nombres. Las dificultades son interpretaciones de lectura, no explicaciones causales comprobadas. No se reajustaron umbrales o etiquetas humanas después de observarlos.')
 # Página9
 page(d);head(d,'7 Interpretabilidad')
 para(d,'Los clásicos usan características explícitas: raíces TF-IDF y longitud. La Figura 4 muestra los seis coeficientes positivos más altos por etiqueta de regresión logística. La magnitud depende de la escala y correlaciones; no es causalidad ni porcentaje de confianza. La longitud también interviene, aunque la figura se limita a términos.')
 image(d,'terminos_clasicos.png','Figura 4 Raíces con mayor coeficiente positivo de regresión logística',6.6)
 para(d,'Los términos asociados a limpieza y ruido apuntan a aseo y sonido; ubicación se apoya en raíces de localización y cercanía; anfitrión en atención y comunicación; precio en valor y costo. La raíz preci ejemplifica el riesgo de colisiones lingüísticas. SVM dispone de su tabla de coeficientes en terminos_clasicos.csv.')
 image(d,'importancia_red.png','Figura 5 Permutation importance de las entradas densas de Keras',5.9)
 para(d,'Para Keras, permutamos por separado cada componente y la longitud en las primeras 600 reseñas de validación, tres veces y con semilla fija. La caída de Macro F1 mide dependencia predictiva. Puede subestimar entradas correlacionadas; no implica que un componente tenga significado lingüístico único. El componente 8 reduce F1 en 0,064 al permutarlo, y el 1 en 0,039. Las barras de error son desviaciones entre permutaciones, no intervalos de población.')
 para(d,'Las cargas absolutas del componente 8 incluyen limpi, orden y comod; las del 1, buen, excelent y ubic. Los dos enfoques usan evidencia léxica semejante, pero la red aprende interacciones en un espacio comprimido. Las cargas solo ayudan a interpretar ese espacio: no equivalen a atribuciones de palabras. La aplicación muestra además contribuciones lineales locales o cambios al quitar términos para la reseña ingresada.')
 # Página10
 page(d);head(d,'7 Análisis de sesgo por idioma')
 para(d,'Elegimos idioma de la reseña como subgrupo relevante. langdetect, con semilla fija, estima el idioma de textos de al menos 20 caracteres; los más cortos se consideran indeterminados. Los textos mixtos pueden quedar asignados al idioma dominante. No se interpreta esta clasificación como identidad del huésped.')
 table(d,'Tabla 12 Comparación por idioma estimado',['Subgrupo','Reseñas','Grupos','Logística','SVM','Keras'],[['Español','67','60','0,931','0,942','0,810'],['Inglés','16','16','0,842','0,914','0,678'],['Indeterminado','16','15','No estimable','No estimable','No estimable'],['Portugués','1','1','No estimable','No estimable','No estimable']], [.95,.8,.8,1.35,1.35,1.35])
 para(d,'La Tabla 12 calcula Macro F1 con las cinco etiquetas. En español e inglés todas tienen positivos, por lo que la comparación no oculta etiquetas ausentes. Los indeterminados no tienen positivos manuales: devolver F1 cero por convención no mediría discriminación útil. Portugués tiene una sola reseña y no permite una estimación. Por eso se muestran como no estimables.')
 para(d,'La diferencia inglés menos español es -0,090 para logística, -0,028 para SVM y -0,132 para Keras. Si se desplegara sin revisión, las menciones de quienes escriben en inglés podrían quedar más frecuentemente sin detectar. No atribuimos esto a nacionalidad ni probamos una disparidad poblacional: 16 reseñas en inglés, la mezcla de temas y la selección dirigida son confusores importantes.')
 head(d,'Implicaciones para operaciones',2)
 para(d,'La app debe usarse como filtro de revisión, no como decisión de sanción, reputación o reembolso. Conviene revisar manualmente textos no detectados, especialmente idiomas con peor resultado o no evaluados. Se necesitan muestras por idioma y alojamiento, anotadas con el mismo criterio, antes de decidir si una diferencia exige otro modelo o más datos.')
 para(d,'No utilizamos barrio o tipo de alojamiento porque no están en el Excel. Incorporarlos sin el archivo listings y sin validar el cruce habría introducido supuestos. El idioma cumple el requisito de un subgrupo relevante usando los datos efectivamente disponibles.')
 # Página11
 page(d);head(d,'8 Despliegue y decisión del modelo')
 para(d,'Aplicación pública: https://utpmachinelearning.streamlit.app/\nCódigo: https://github.com/nickymarie090490/utpmachinelearnig')
 para(d,'Desplegamos regresión logística como selección inicial y modelo recomendado. SVM tiene 0,021 puntos más de Macro F1 y 80% de coincidencia completa frente a 74%, pero carece de probabilidades sin un paso adicional de calibración. La logística ya cumple la salida probabilística solicitada y permite explicar directamente las contribuciones de cada raíz. Su inferencia de 0,064 ms por reseña en lote es ligeramente menor que 0,067 de SVM; no justificamos la elección por una diferencia de tiempo grande.')
 para(d,'Keras cuesta 25,809 s de búsqueda y entrenamiento frente a 14,206 de logística, ocupa 12,128 MB frente a 0,297 MB y tarda 0,101 ms por reseña frente a 0,064. Su F1 inferior y dependencias adicionales no justifican elegirla para producción. Los tres quedan disponibles en la app para demostrar la comparación académica.')
 head(d,'Pasos para reproducir el despliegue',2)
 for txt in ['1 Instalar Python 3.12 y las versiones de requirements.txt en un entorno virtual.', '2 Publicar app.py, los módulos Python, modelos, manifiesto y resultados agregados en el repositorio.', '3 Conectar el repositorio en Streamlit Community Cloud, servicio gratuito; seleccionar main, app.py y Python 3.12.', '4 Instalar dependencias y cargar el modelo guardado. La app no entrena al abrirse y no necesita API de pago.', '5 Verificar una reseña, un CSV UTF-8 con comments, selector de modelos, explicación local y la pestaña Red Keras.', '6 Mantener código, resultados y modelo de la misma versión en un commit; los hashes verifican integridad.']:
  para(d,txt)
 head(d,'Explicación de la predicción y guía de uso',2)
 para(d,'La entrada individual acepta hasta 15.000 caracteres y el CSV hasta 5.000 filas no vacías. Se muestran cinco temas, puntuaciones cuando existen y una explicación para el caso ingresado. En clásicos se presenta el producto del valor de la característica por su coeficiente; en Keras el cambio al quitar términos. Ambos métodos se explican como aproximaciones predictivas, no causalidad. La guía de uso aclara que detectar un tema no clasifica sentimiento.')
 para(d,'En la regresión del fallo observado, hay mucho ruido ahora devuelve solo ruido en los tres modelos. En la versión anterior Keras daba aproximadamente 3,5%; la nueva salida es muy alta, pero no calibrada. Ese caso aislado no borra los 16 falsos negativos de ruido del test. Las pruebas de entradas vacías y sin letras generan mensajes claros en lugar de inferencias confusas.')
 # Página12
 page(d);head(d,'9 Diagrama del pipeline del proyecto')
 image(d,'pipeline.png','Figura 6 Flujo desde el archivo original hasta los dos niveles y la app',6.6)
 para(d,'La Figura 6 marca la partición antes de ajustar transformaciones. Los dos niveles utilizan la misma información original y la misma separación de alojamientos. La reserva humana se evalúa tras seleccionar modelos; la limitación por conocimiento previo queda declarada. La normalización para excluir copias no ajusta parámetros estadísticos: vocabulario, IDF, SVD y escalas permanecen dentro de la preparación del pipeline.')
 head(d,'Verificación mediante pruebas unitarias',2)
 para(d,f"Se ejecutaron {tests['pruebas']} pruebas con pytest: {tests['fallos']} fallos, {tests['errores']} errores y {tests['omitidas']} omitidas, en {tests['segundos']:.2f} s. Verifican normalización, stemming, reglas corregidas, entradas inválidas, vocabulario sin fuga, salida densa con longitud, separación de textos y grupos, F1 y matrices conocidas, inferencia, explicación local, integridad ante corrupción, arquitectura, early stopping configurado y curvas por entidades.")
 para(d,'AppTest comprobó análisis con los tres modelos y rechazo de una entrada sin letras, sin excepciones. El fallo de ruido se agregó como prueba de regresión para logística y Keras. Una prueba que pasa demuestra el contrato concreto que comprueba; no convierte toda predicción en correcta ni reemplaza la evaluación humana. Los resultados de pytest y los archivos agregados quedan en datos para trazabilidad.')
 # Página13
 page(d);head(d,'Reflexión crítica')
 head(d,'1 Preparación de los datos',2)
 para(d,'El mayor trabajo fue construir un target consistente a partir de texto y separar alojamientos antes de entrenar. Se quitaron HTML, puntuación, acentos y stopwords, se aplicó stemming y se añadieron longitud y TF-IDF. Se excluyeron 54.100 duplicados normalizados, 601 textos sin letras y los alojamientos o copias del test. Conservar reseñas múltiples habría inflado la idea de cantidad independiente.')
 para(d,'Elegimos menciones porque operaciones quiere saber qué temas revisar. Macro F1 evita que ubicación domine ruido y precio. La falta de acuerdo entre anotadores y la supervisión débil cuestan más que una simple limpieza: una métrica automática alta puede significar aprender bien reglas imperfectas. La Tabla 3 documenta qué se descartó, en lugar de atribuir toda reducción a limpieza.')
 head(d,'2 Los dos enfoques',2)
 para(d,'La red no mejoró frente al clásico final: perdió 0,132 puntos de Macro F1 respecto a logística y 0,153 frente a SVM. Mejoró de 0,635 a 0,781 frente a su versión previa, pero es una comparación de versiones con cambios conjuntos y un test ya conocido, no una estimación ciega de mejora causal. El costo adicional de entrenamiento, inferencia y almacenamiento no se justifica para la aplicación elegida.')
 para(d,'El clásico depende de ingeniería explícita: limpieza, raíces, TF-IDF y longitud. La red también recibe esas características, comprimidas por SVD, y aprende combinaciones no lineales. El espacio denso no garantiza comprensión de sinónimos. Las Figuras 4 y 5 muestran evidencia compartida, mientras que la Tabla 11 revela errores por contexto y compresión.')
 head(d,'3 La decisión de despliegue',2)
 para(d,'Pondríamos logística en producción como apoyo de revisión. Aceptamos un F1 0,021 inferior a SVM para disponer de puntuaciones por tema y contribuciones interpretables sin añadir calibración de otro modelo. Los dos tamaños en disco y tiempos de inferencia son casi iguales; la decisión descansa principalmente en funcionalidad y mantenimiento, no en ahorro considerable. Keras permanece como comparación obligatoria, no como recomendación operacional.')
 para(d,'Si operaciones prefiriera decisiones binarias y maximizar coincidencia con las etiquetas de este estudio, SVM sería una alternativa defendible. La app muestra ambos para que se entienda el costo de la decisión. No se afirma que sus porcentajes sean probabilidades calibradas ni que una detección autorice una acción automática.')
 # Página14
 page(d);head(d,'Reflexión crítica y límites del estudio')
 head(d,'4 Evaluación de resultados',2)
 para(d,'Logística falla en 26 reseñas completas, especialmente por omisiones de ubicación y falsas asociaciones de precio; precio tiene F1 0,818 y es su categoría más débil. La Tabla 11 muestra giros indirectos, múltiples temas y colisiones de raíces. En ruido, su recall es 0,960 frente a 0,360 de Keras. Ambos enfoques pueden ejecutar sin errores técnicos y aun así equivocarse semánticamente.')
 para(d,'La partición ingenua no produjo una ventaja de F1: logística quedó 0,002794 por debajo de la agrupada y SVM prácticamente igual. Aun así, compartió 2.690 alojamientos por pliegue. Esta diferencia enseña que la validez depende de la población objetivo y de la separación, no de que una tabla arroje la dirección esperada.')
 head(d,'5 Reflexión ética y uso responsable',2)
 para(d,'Usar estas detecciones como conteo de quejas podría perjudicar a anfitriones que recibieron elogios o menciones neutrales. Los falsos negativos pueden dejar sin revisión problemas de huéspedes; los falsos positivos pueden desviar recursos o dañar reputaciones. Se deben revisar los textos antes de sanciones o reembolsos. Se descartaron nombres e identificadores de reseñadores del modelo y no se publican reseñas originales.')
 para(d,'En logística, inglés obtuvo 0,842 y español 0,931. Una revisión automática sin control podría omitir más menciones en inglés. El subgrupo inglés es pequeño y la muestra dirigida, por lo que no acusamos una discriminación poblacional comprobada. Portugués y textos indeterminados no tienen evidencia suficiente para evaluar las cinco categorías.')
 head(d,'6 Limitaciones del estudio',2)
 para(d,'Las etiquetas de entrenamiento son reglas, la muestra manual tiene solo 100 reseñas y 90 alojamientos, y fue consultada antes. No hay test humano nuevo, acuerdo entre anotadores, fecha acreditada del dump ni validación externa. Los clásicos siguen mejorando al 100% y Keras es inestable; los datos no permiten declarar suficiencia. Una semilla y CPU limitada reducen la evidencia sobre variabilidad. La mayor deficiencia restante es evaluativa, no de ejecución del código.')
 para(d,'Con tres meses y una GPU dedicaríamos primero esfuerzo a un test ciego por alojamiento e idioma y a etiquetado humano con doble revisión. Después repetiríamos semillas, calibraríamos puntuaciones con validación independiente y evaluaríamos más tamaños de representación. Las mejoras deben responder a errores medidos y a necesidades de operaciones; una arquitectura más grande no sustituye mejor supervisión.')
 head(d,'Referencias',2)
 para(d,'Montenegro J. Proyecto Final del Curso de Machine Learning Grupo A, guía suministrada, pp. 4-7.\nInside Airbnb. Get the Data. https://insideairbnb.com/get-the-data/\nscikit-learn. GroupKFold, TruncatedSVD y permutation importance. https://scikit-learn.org/stable/\nKeras. EarlyStopping y Dropout. https://keras.io/api/\nInsumos del equipo: Etiquetados.xlsx y Problema_2.ipynb; versiones identificadas en origen.json y bitácora.')
 # Apéndice completo fuera del límite principal.
 page(d);head(d,'Apéndice Código fuente')
 para(d,'El informe principal termina antes de este apéndice. Se reproduce en texto el código de preparación, entrenamiento, evaluación, inferencia, figuras, app y pruebas. El repositorio y el paquete incluyen los archivos ejecutables y las versiones exactas. El apéndice no cuenta para el límite de 5 a 15 páginas de la guía.')
 for f in ['texto.py','preparar_excel.py','entrenar.py','evaluar.py','keras_red.py','inferencia.py','generar_informe.py','app.py','tests/test_proyecto.py']:
  head(d,'Código '+f.replace('.py','').replace('_',' ').replace('/',' '),2)
  content=(P/f).read_text()
  # Cada línea se preserva; Word realiza el ajuste visual sin cortar el código.
  p=d.add_paragraph();p.paragraph_format.line_spacing=1;p.paragraph_format.space_after=Pt(6)
  run=p.add_run(content);run.font.name='Consolas';run.font.size=Pt(7.5)
 d.save(O/'Informe_Proyecto_Equipo6.docx');print('Informe Word creado')

def bitacora():
 d=base();d.add_paragraph('Bitácora del proyecto de clasificación de reseñas','Title')
 para(d,'Equipo 6 · Problema 2 · Revisión del 4 de octubre de 2026\nRegistro reconstruido de decisiones observadas en el código, archivos y ejecución; no acredita una fecha del dump ni actividades no documentadas.')
 head(d,'1 Preparación y decisiones previas')
 para(d,'4 octubre 2026. Se recibió Etiquetados.xlsx con 690.112 reseñas y 100 etiquetas humanas completas. El notebook identifica Santiago. Se mantuvieron las cinco categorías de menciones, Macro F1 como métrica principal y separación por listing_id. La longitud se combinó con TF-IDF mediante ColumnTransformer. Se excluyeron los alojamientos humanos y copias de sus textos antes del muestreo de 25.000.')
 para(d,'4 octubre 2026. Se conservaron logística y SVM como comparadores clásicos. La implementación anterior de MLPClassifier no cumplía Keras obligatorio y se sustituyó por tf.keras con TruncatedSVD, dropout y early stopping. La primera versión Keras publicada alcanzó Macro F1 manual 0,635, frente a 0,912 de logística y 0,926 de SVM; sus curvas y configuración se guardaron.')
 para(d,'4 octubre 2026. La prueba pública hay mucho ruido se procesó sin excepción, pero la red asignó aproximadamente 3,5% a ruido y no lo detectó. Se declaró el fallo; no se forzó una etiqueta en la app. La revisión posterior comprobó que ruido sí estaba en el vocabulario, pero los 64 componentes conservaban poca señal, y la etiqueta representaba cerca del 4% del entrenamiento.')
 head(d,'Trazabilidad de la versión de datos',2)
 para(d,'Archivo Etiquetados.xlsx. SHA-256 b6b561fc6334a9b2195a45e4dacebd5804983c3c7baa4a7d22cafce49a179e4a. Período de reseñas 13 noviembre 2010 a 1 julio 2026. No consta la fecha de descarga del dump: queda pendiente obtenerla de la fuente del equipo. No se presenta la fecha del Excel como fecha Inside Airbnb.')
 para(d,'La muestra humana ya había sido observada, por lo que se registró su límite de evaluación ciega. Los cambios finales se seleccionaron con validación automática sin ajustar umbrales humanos. Es necesaria una nueva muestra manual independiente para cerrar esa condición estricta de la guía.')
 page(d);head(d,'2 Revisión según la guía y alternativas ensayadas')
 para(d,'4 octubre 2026. Se auditó la guía completa. Se añadieron comparación ingenua y agrupada, curva de aprendizaje por entidades, costos de búsqueda e inferencia, tamaños, matrices multietiqueta, cinco errores, coeficientes, permutation importance, idiomas, pruebas y explicación local. Se corrigió la regla cost para no confundir Costanera con precio y se añadieron loud y loudly. Se limitó el vocabulario a 6.000 para acotar recursos.')
 table(d,'Tabla 1 Alternativas ensayadas y motivo del descarte',['Alternativa','Evidencia medida','Decisión'],[
 ['MLP 64 sin ponderar','F1 automático 0,639; pierde señal de etiquetas escasas','Descartado como candidato final'],['MLP 128 con BCE ponderada','F1 automático 0,728; dos capas 128/64','Descartado frente al 0,781 de256'],['Bigramas en clásicos','GridSearchCV incluyó unigramas y bigramas; ganó unigramas','Descartados en configuración final'],['KFold por filas','Comparte 2.690 alojamientos por pliegue; F1 no aumenta','Descartado como validador final'],['Red sin dropout ni early stopping','BCE validación 0,303 frente a0,109; F1 0,783 frente a0,781','Descartada por pérdida y sobreajuste'],['SVM para app por defecto','Mayor F1 0,934, sin salida probabilística directa','Conservado para comparar; elegida logística']], [2.0,2.8,1.8])
 para(d,'Se seleccionó SVD256, capas128/64, dropout0,20, Adam0,001, batch128 y pérdida ponderada por etiqueta. Los pesos se calcularon solo con entrenamiento. EarlyStopping ejecutó12 de100 épocas y restauró6. No se atribuyó la ganancia a compresión o ponderación por separado, porque ambos cambios se probaron conjuntamente.')
 para(d,'La ablación mantuvo arquitectura, representación, pesos, semilla y optimizador; solo quitó dropout y parada. La BCE de validación bajó63,9% con controles, aunque el F1 automático fue ligeramente menor. Esta diferencia se mantuvo en el informe; no se ocultó para sostener la efectividad.')
 page(d);head(d,'3 Evaluación final y publicación')
 para(d,'4 octubre 2026. Se completaron cinco tamaños de entrenamiento por alojamiento. Logística y SVM siguieron ascendiendo al100%; Keras fue inestable. Se declaró que no hay evidencia de suficiencia. La comparación agrupada evitó alojamientos compartidos, aun cuando el F1 ingenuo resultó similar o menor.')
 para(d,'La evaluación final mantuvo100 reseñas de90 alojamientos. Resultados Macro F1: logística0,913, SVM0,934 y Keras0,781. En ruido, Keras acertó9 de25 positivos; no se interpretó el éxito de una frase como solución general. Se analizaron cinco errores y diferencias entre español67 reseñas e inglés16; los subgrupos sin positivos se reportaron como no estimables.')
 para(d,'Se ejecutaron25 pruebas unitarias sin fallos, errores ni omisiones. AppTest verificó los tres modelos y entradas sin letras. Se agregó la regresión hay mucho ruido y se confirmó detección en los tres. Las pruebas de integridad rechazan un preparador corrupto. Los costos se midieron tras calentamiento y se distinguieron de carga inicial.')
 para(d,'La decisión fue desplegar regresión logística por probabilidades no calibradas, explicación lineal y mantenimiento simple. SVM y Keras permanecen disponibles para la comparación. Se generaron el informe Word, PDF con código en apéndice, README.txt, requirements con versiones exactas, bitácora y paquete reproducible. El repositorio contiene código, modelos y resultados agregados; los datos originales y textos del test no se publican.')
 head(d,'Pendientes que requieren evidencia del equipo',2)
 para(d,'Acreditar la ciudad asignada y la fecha del dump con el registro de descarga. Construir una nueva muestra humana de alojamientos no usados ni revisados para un test final ciego. Confirmar el criterio de anotación y, de ser posible, acuerdo entre anotadores. Preparar la exposición oral de12 minutos y la demostración; no se añadió una presentación porque la guía la declara opcional.')
 head(d,'Reproducción y control de versiones',2)
 para(d,'Semilla20261004; Python3.12.14; TensorFlow2.21.0 y Keras3.15.1. El comando preparar_excel.py reconstruye las muestras desde el Excel, entrenar.py reproduce modelos y análisis, pytest verifica contratos y streamlit run app.py inicia la aplicación. Los hashes del preparador y modelo verifican que las piezas corresponden a la misma versión. Pequeños redondeos y tiempos pueden variar por plataforma.')
 d.save(O/'Bitacora_Proyecto_Equipo6.docx');print('Bitácora Word creada')
if __name__=='__main__':informe();bitacora()
